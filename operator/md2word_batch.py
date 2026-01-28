"""Markdown 转 Word 工具 - 支持单文件和多文件输入

将 Markdown 文件（.md）转换为 Word 文档（.docx）。
支持多种转换后端：
1. pypandoc (推荐，需要安装 Pandoc)
2. comtypes + Word (Windows，需要安装 Microsoft Word，通过 HTML 中转)
3. python-docx (纯Python方案，仅保留基础文本结构)
"""
import os
import tempfile
import uuid
import shutil
import requests
from typing import Optional, List, Dict

from langchain_core.tools import tool

# MCP 服务地址（用于注册文件下载）
MCP_API_BASE = os.environ.get("MCP_API_BASE", "http://localhost:7000/api/mcp")

# 支持的 Markdown 文件扩展名
MD_EXTENSIONS = ('.md', '.markdown')


def register_file_download(file_path: str, expires_in: int = 3600) -> Optional[str]:
    """向 MCP 服务注册文件，获取下载 URL"""
    try:
        resp = requests.post(
            f"{MCP_API_BASE}/files/register",
            json={"file_path": file_path, "expires_in": expires_in},
            timeout=10
        )
        if resp.status_code == 200:
            return resp.json().get("download_url")
    except Exception:
        pass
    return None


def _validate_md_file(path: str) -> Optional[str]:
    """验证并返回有效的 Markdown 文件路径"""
    if os.path.isfile(path) and path.lower().endswith(MD_EXTENSIONS):
        return os.path.abspath(path)
    return None


def _convert_with_pypandoc(md_path: str, output_path: str) -> tuple:
    """使用 pypandoc 转换（最佳质量）"""
    try:
        import pypandoc
        # 检查 pandoc 是否存在
        try:
            pypandoc.get_pandoc_version()
        except OSError:
             return False, "未找到 Pandoc，请安装 Pandoc 或尝试其他转换方式"

        output = pypandoc.convert_file(md_path, 'docx', outputfile=output_path)
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成 .docx 文件"
            
    except ImportError:
        return False, "未安装 pypandoc，请执行: pip install pypandoc"
    except Exception as e:
        return False, f"pypandoc 转换异常: {str(e)}"


def _convert_with_comtypes(md_path: str, output_path: str) -> tuple:
    """使用 comtypes + Word 转换（Windows，通过 HTML 中转）"""
    html_path = None
    try:
        import markdown
        import comtypes.client
        
        # 1. MD -> HTML
        with open(md_path, 'r', encoding='utf-8') as f:
            text = f.read()
        
        html = markdown.markdown(text, extensions=['tables', 'fenced_code'])
        
        # 包装完整的 HTML 结构以确保编码正确
        full_html = f"""
        <html>
        <head>
            <meta charset="utf-8">
        </head>
        <body>
            {html}
        </body>
        </html>
        """
        
        # 保存临时 HTML 文件
        base_name = os.path.splitext(os.path.basename(md_path))[0]
        html_path = os.path.join(tempfile.gettempdir(), f"{base_name}_{uuid.uuid4().hex[:8]}.html")
        with open(html_path, 'w', encoding='utf-8') as f:
            f.write(full_html)
            
        # 2. HTML -> Word (via COM)
        word = comtypes.client.CreateObject("Word.Application")
        word.Visible = False
        
        try:
            doc = word.Documents.Open(html_path)
            # 16 = wdFormatDocumentDefault (docx)
            doc.SaveAs(output_path, FileFormat=16)
            doc.Close()
        finally:
            word.Quit()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成 .docx 文件"
            
    except ImportError as e:
        return False, f"缺少依赖库: {str(e)} (需要 markdown 和 comtypes)"
    except Exception as e:
        error_msg = str(e)
        if "Word.Application" in error_msg:
            return False, "未安装 Microsoft Word 或无法启动"
        return False, f"comtypes 转换异常: {error_msg}"
    finally:
        if html_path and os.path.exists(html_path):
            try:
                os.remove(html_path)
            except:
                pass


def _convert_with_python(md_path: str, output_path: str) -> tuple:
    """使用纯Python库转换（基础转换，格式丢失较多）"""
    try:
        from docx import Document
        
        doc = Document()
        
        with open(md_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            else:
                doc.add_paragraph(line)
        
        doc.save(output_path)
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成 .docx 文件"
            
    except ImportError as e:
        return False, f"缺少依赖库: {str(e)}，请执行: pip install python-docx"
    except Exception as e:
        return False, f"Python 转换异常: {str(e)}"


def _detect_best_converter() -> str:
    """检测最佳的转换方案"""
    # 1. 优先检查 pypandoc (Pandoc)
    try:
        import pypandoc
        try:
            pypandoc.get_pandoc_version()
            return 'pypandoc'
        except OSError:
            pass
    except ImportError:
        pass
    
    # 2. 检查 comtypes + markdown (Windows)
    try:
        import comtypes.client
        import markdown
        try:
            word = comtypes.client.CreateObject("Word.Application")
            word.Quit()
            return 'comtypes'
        except:
            pass
    except ImportError:
        pass
    
    # 3. 最后使用纯 Python 方案
    try:
        import docx
        return 'python'
    except ImportError:
        pass
    
    return 'none'


def _convert_single_md(
    md_path: str,
    output_path: Optional[str] = None,
    converter: str = "auto",
) -> dict:
    """转换单个 Markdown 文件为 Word（内部函数）"""
    
    if not os.path.isfile(md_path):
        return {
            "success": False,
            "source_file": md_path,
            "word_path": None,
            "download_url": None,
            "message": f"错误：文件不存在 - {md_path}"
        }
    
    ext = os.path.splitext(md_path)[1].lower()
    if ext not in MD_EXTENSIONS:
        return {
            "success": False,
            "source_file": md_path,
            "word_path": None,
            "download_url": None,
            "message": f"错误：不支持的文件格式 '{ext}'，仅支持 .md"
        }
    
    # 确定输出路径
    if output_path is None:
        base_name = os.path.splitext(os.path.basename(md_path))[0]
        output_path = os.path.join(
            tempfile.gettempdir(),
            f"{base_name}_{uuid.uuid4().hex[:8]}.docx"
        )
    
    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 选择转换器
    actual_converter = converter
    if converter == 'auto':
        actual_converter = _detect_best_converter()
    
    if actual_converter == 'none':
        return {
            "success": False,
            "source_file": md_path,
            "word_path": None,
            "download_url": None,
            "message": "错误：未找到可用的转换器 (需安装 pandoc, markdown+comtypes 或 python-docx)"
        }
    
    # 执行转换
    if actual_converter == 'pypandoc':
        success, error = _convert_with_pypandoc(md_path, output_path)
    elif actual_converter == 'comtypes':
        success, error = _convert_with_comtypes(md_path, output_path)
    elif actual_converter == 'python':
        success, error = _convert_with_python(md_path, output_path)
    else:
        return {
            "success": False,
            "source_file": md_path,
            "word_path": None,
            "download_url": None,
            "message": f"错误：未知的转换器 '{actual_converter}'"
        }
    
    if not success:
        return {
            "success": False,
            "source_file": md_path,
            "word_path": None,
            "download_url": None,
            "message": f"转换失败: {error}",
            "converter_used": actual_converter
        }
    
    download_url = register_file_download(output_path)
    
    result = {
        "success": True,
        "source_file": md_path,
        "word_path": output_path,
        "download_url": download_url,
        "message": f"转换成功：{os.path.basename(md_path)} -> {os.path.basename(output_path)}",
        "converter_used": actual_converter,
        "file_size_kb": round(os.path.getsize(output_path) / 1024, 2)
    }
    
    return result


@tool
def md_to_word(
    input_path: str,
    converter: str = "auto",
    output_dir: Optional[str] = None,
) -> dict:
    """将 Markdown 文件转换为 Word (.docx) 格式。
    
    支持单个或多个文件输入。

    Args:
        input_path: 输入文件路径，支持：
            - 单个文件："/path/to/file.md"
            - 多个文件（逗号分隔）："/path/file1.md, /path/file2.md"
        converter: 转换器，可选 'pypandoc'/'comtypes'/'python'/'auto'（默认自动检测）。
        output_dir: 输出目录，若不指定则使用临时目录。

    Returns:
        包含转换结果的字典。
    """
    if not input_path or not input_path.strip():
        return {
            "success": False,
            "total_files": 0,
            "processed": 0,
            "failed": 0,
            "results": [],
            "message": "错误：请提供有效的输入路径"
        }
    
    # 解析输入路径（支持逗号分隔）
    input_paths = [p.strip() for p in input_path.split(',') if p.strip()]
    
    # 验证并收集所有 MD 文件
    all_md_files = []
    invalid_paths = []
    for path in input_paths:
        valid_path = _validate_md_file(path)
        if valid_path:
            all_md_files.append(valid_path)
        else:
            invalid_paths.append(path)
    
    all_md_files = sorted(set(all_md_files))
    
    if not all_md_files:
        error_msg = f"错误：未找到有效的 Markdown 文件（输入：{input_path}）"
        if invalid_paths:
            error_msg += f"\n无效路径: {', '.join(invalid_paths)}"
        return {
            "success": False,
            "total_files": 0,
            "processed": 0,
            "failed": 0,
            "results": [],
            "message": error_msg
        }
    
    # 确定输出目录
    if output_dir is None:
        output_dir = os.path.join(tempfile.gettempdir(), f"md2word_{uuid.uuid4().hex[:8]}")
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 检测转换器（只检测一次）
    actual_converter = converter
    if converter == 'auto':
        actual_converter = _detect_best_converter()
    
    # 处理所有文件
    results = []
    processed = 0
    failed = 0
    
    for md_file in all_md_files:
        base_name = os.path.splitext(os.path.basename(md_file))[0]
        output_path = os.path.join(output_dir, f"{base_name}.docx")
        
        # 避免文件名冲突
        counter = 1
        while os.path.exists(output_path):
            output_path = os.path.join(output_dir, f"{base_name}_{counter}.docx")
            counter += 1
        
        result = _convert_single_md(
            md_path=md_file,
            output_path=output_path,
            converter=actual_converter,
        )
        
        results.append(result)
        
        if result["success"]:
            processed += 1
        else:
            failed += 1
    
    # 构建返回结果
    success = failed == 0 and processed > 0
    
    # 构建文件映射列表
    file_mappings = []
    for r in results:
        if r["success"]:
            file_mappings.append({
                "source": r["source_file"],
                "output": r["word_path"],
                "download_url": r["download_url"],
            })
    
    response = {
        "success": success,
        "total_files": len(all_md_files),
        "processed": processed,
        "failed": failed,
        "results": results,
        "file_mappings": file_mappings,
        "output_dir": output_dir,
        "converter_used": actual_converter,
        "message": f"转换完成：共 {len(all_md_files)} 个文件，成功 {processed} 个，失败 {failed} 个"
    }
    
    # 单文件时返回简化结果
    if len(all_md_files) == 1:
        single = results[0]
        response.update({
            "word_path": single.get("word_path"),
            "download_url": single.get("download_url"),
            "file_size_kb": single.get("file_size_kb"),
            "message": single.get("message"),
        })
    
    return response


if __name__ == "__main__":
    # 检测环境
    print("=== 检测转换器 ===")
    converter = _detect_best_converter()
    print(f"可用转换器: {converter}")
    
    # 创建测试文件
    test_md_content = """# 测试文档
## 章节 1
这是一个**测试** Markdown 文档。
- 列表项 1
- 列表项 2

### 表格测试
| 姓名 | 年龄 |
| --- | --- |
| 张三 | 25 |
| 李四 | 30 |
"""
    test_dir = os.path.join(tempfile.gettempdir(), "md2word_test")
    if not os.path.exists(test_dir):
        os.makedirs(test_dir)
        
    test_file = os.path.join(test_dir, "test.md")
    with open(test_file, "w", encoding="utf-8") as f:
        f.write(test_md_content)
        
    print(f"\n=== 测试单文件: {test_file} ===")
    result = md_to_word.invoke({"input_path": test_file})
    print(f"成功: {result['success']}")
    print(f"消息: {result['message']}")
    print(f"下载链接: {result.get('download_url')}")
    print(f"转换器: {result.get('converter_used')}")
    
    # 清理
    # shutil.rmtree(test_dir)
