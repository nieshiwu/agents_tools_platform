"""PPT/PPTX 转 PDF 工具 - 支持单文件和多文件输入

将 PowerPoint 演示文稿转换为 PDF 格式。
支持多种转换后端：
1. LibreOffice (Linux/Docker 推荐，免费开源)
2. comtypes + PowerPoint (Windows，需要安装 Microsoft PowerPoint)
3. python-pptx + reportlab (纯Python方案，基础转换)
"""
import os
import subprocess
import tempfile
import uuid
import shutil
import requests
from typing import Optional, List, Dict

from langchain_core.tools import tool

# MCP 服务地址（用于注册文件下载）
MCP_API_BASE = os.environ.get("MCP_API_BASE", "http://localhost:7000/api/mcp")

# 支持的 PPT 文件扩展名
PPT_EXTENSIONS = ('.ppt', '.pptx')


def register_file_download(file_path: str, expires_in: int = 3600) -> Optional[str]:
    """向 MCP 服务注册文件，获取下载 URL"""
    try:
        resp = requests.post(
            f"{MCP_API_BASE}/files/register",
            json={"file_path": file_path, "expires_in": expires_in},
            timeout=10
        )
        if resp.status_code == 200:
            return resp.json().get("download_url")
    except Exception:
        pass
    return None


def _validate_ppt_file(path: str) -> Optional[str]:
    """验证并返回有效的 PPT 文件路径
    
    Args:
        path: 文件路径
        
    Returns:
        有效的绝对路径，如果文件无效则返回 None
    """
    if os.path.isfile(path) and path.lower().endswith(PPT_EXTENSIONS):
        return os.path.abspath(path)
    return None


def _convert_with_libreoffice(ppt_path: str, output_path: str) -> tuple:
    """使用 LibreOffice 转换单个文件（Linux/Docker 推荐）"""
    try:
        output_dir = os.path.dirname(output_path) or tempfile.gettempdir()
        
        cmd = [
            'soffice',
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
            ppt_path
        ]
        
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=180
        )
        
        if result.returncode != 0:
            return False, f"LibreOffice 转换失败: {result.stderr}"
        
        base_name = os.path.splitext(os.path.basename(ppt_path))[0]
        generated_pdf = os.path.join(output_dir, f"{base_name}.pdf")
        
        if generated_pdf != output_path and os.path.exists(generated_pdf):
            shutil.move(generated_pdf, output_path)
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except FileNotFoundError:
        return False, "未找到 LibreOffice，请安装: apt-get install libreoffice"
    except subprocess.TimeoutExpired:
        return False, "转换超时，PPT文件可能过大或过于复杂"
    except Exception as e:
        return False, f"LibreOffice 转换异常: {str(e)}"


def _convert_batch_with_libreoffice(ppt_files: List[str], output_dir: str) -> Dict[str, tuple]:
    """使用 LibreOffice 批量转换多个文件（只启动一次，大幅提升速度）
    
    Args:
        ppt_files: PPT 文件路径列表
        output_dir: 输出目录
        
    Returns:
        字典 {ppt_path: (success, error_or_pdf_path)}
    """
    results = {}
    
    if not ppt_files:
        return results
    
    try:
        # 批量转换：一次性传入所有文件，只启动一次 LibreOffice
        cmd = [
            'soffice',
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
        ] + ppt_files
        
        # 根据文件数量动态调整超时时间（每个文件约60秒）
        timeout = max(180, len(ppt_files) * 60)
        
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=timeout
        )
        
        # 检查每个文件的转换结果
        for ppt_path in ppt_files:
            base_name = os.path.splitext(os.path.basename(ppt_path))[0]
            generated_pdf = os.path.join(output_dir, f"{base_name}.pdf")
            
            if os.path.exists(generated_pdf):
                results[ppt_path] = (True, generated_pdf)
            else:
                # 单个文件转换失败
                results[ppt_path] = (False, f"转换失败，未生成 PDF（可能是文件损坏或格式不支持）")
        
        # 如果整体命令失败且没有任何成功的文件
        if result.returncode != 0 and not any(r[0] for r in results.values()):
            error_msg = f"LibreOffice 批量转换失败: {result.stderr}"
            for ppt_path in ppt_files:
                if ppt_path not in results or not results[ppt_path][0]:
                    results[ppt_path] = (False, error_msg)
                    
    except FileNotFoundError:
        error = "未找到 LibreOffice，请安装: apt-get install libreoffice"
        for ppt_path in ppt_files:
            results[ppt_path] = (False, error)
    except subprocess.TimeoutExpired:
        error = "批量转换超时，文件可能过多或过大"
        for ppt_path in ppt_files:
            if ppt_path not in results:
                results[ppt_path] = (False, error)
    except Exception as e:
        error = f"LibreOffice 批量转换异常: {str(e)}"
        for ppt_path in ppt_files:
            results[ppt_path] = (False, error)
    
    return results


def _convert_with_comtypes(ppt_path: str, output_path: str) -> tuple:
    """使用 comtypes + PowerPoint 转换（Windows）"""
    try:
        import comtypes.client
        
        powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
        powerpoint.Visible = 1
        
        try:
            presentation = powerpoint.Presentations.Open(ppt_path)
            presentation.SaveAs(output_path, 32)  # 32 = PDF
            presentation.Close()
        finally:
            powerpoint.Quit()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except ImportError:
        return False, "未安装 comtypes，请执行: pip install comtypes"
    except Exception as e:
        error_msg = str(e)
        if "Powerpoint.Application" in error_msg:
            return False, "未安装 Microsoft PowerPoint 或无法启动"
        return False, f"comtypes 转换异常: {error_msg}"


def _convert_with_python(ppt_path: str, output_path: str) -> tuple:
    """使用纯Python库转换（基础转换）"""
    try:
        from pptx import Presentation
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from PIL import Image
        import io
        
        font_name = 'Helvetica'
        try:
            chinese_fonts = [
                '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
                '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
                'C:/Windows/Fonts/simhei.ttf',
                'C:/Windows/Fonts/msyh.ttc',
            ]
            for font_path in chinese_fonts:
                if os.path.exists(font_path):
                    pdfmetrics.registerFont(TTFont('ChineseFont', font_path))
                    font_name = 'ChineseFont'
                    break
        except:
            pass
        
        prs = Presentation(ppt_path)
        slide_width = prs.slide_width.inches
        slide_height = prs.slide_height.inches
        page_size = (slide_width * 72, slide_height * 72)
        c = canvas.Canvas(output_path, pagesize=page_size)
        
        for slide_num, slide in enumerate(prs.slides, 1):
            c.setFillColorRGB(1, 1, 1)
            c.rect(0, 0, page_size[0], page_size[1], fill=1)
            
            y_position = page_size[1] - 50
            
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            try:
                                c.setFont(font_name, 12)
                                c.setFillColorRGB(0, 0, 0)
                                
                                max_width = page_size[0] - 100
                                if c.stringWidth(text, font_name, 12) > max_width:
                                    while c.stringWidth(text + "...", font_name, 12) > max_width and len(text) > 10:
                                        text = text[:-1]
                                    text = text + "..."
                                
                                c.drawString(50, y_position, text)
                                y_position -= 20
                                
                                if y_position < 50:
                                    break
                            except:
                                pass
                
                if shape.shape_type == 13:  # Picture
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        
                        img = Image.open(io.BytesIO(image_bytes))
                        img_reader = ImageReader(img)
                        
                        img_width = min(shape.width.inches * 72, page_size[0] - 100)
                        img_height = min(shape.height.inches * 72, page_size[1] - 100)
                        
                        x_pos = (page_size[0] - img_width) / 2
                        y_pos = max(50, y_position - img_height)
                        
                        c.drawImage(img_reader, x_pos, y_pos, 
                                   width=img_width, height=img_height,
                                   preserveAspectRatio=True)
                        
                        y_position = y_pos - 20
                    except:
                        pass
            
            c.setFont(font_name, 10)
            c.setFillColorRGB(0.5, 0.5, 0.5)
            c.drawCentredString(page_size[0] / 2, 20, f"第 {slide_num} 页")
            c.showPage()
        
        c.save()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except ImportError as e:
        return False, f"缺少依赖库: {str(e)}，请执行: pip install python-pptx reportlab Pillow"
    except Exception as e:
        return False, f"Python 转换异常: {str(e)}"


def _detect_best_converter() -> str:
    """检测最佳的转换方案"""
    try:
        result = subprocess.run(['soffice', '--version'], capture_output=True, timeout=5)
        if result.returncode == 0:
            return 'libreoffice'
    except:
        pass
    
    try:
        import comtypes.client
        try:
            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.Quit()
            return 'comtypes'
        except:
            pass
    except ImportError:
        pass
    
    try:
        import pptx
        import reportlab
        return 'python'
    except ImportError:
        pass
    
    return 'none'


def _convert_single_ppt(
    ppt_path: str,
    output_path: Optional[str] = None,
    converter: str = "auto",
) -> dict:
    """转换单个 PPT 文件为 PDF（内部函数）"""
    
    if not os.path.isfile(ppt_path):
        return {
            "success": False,
            "source_file": ppt_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：文件不存在 - {ppt_path}"
        }
    
    ext = os.path.splitext(ppt_path)[1].lower()
    if ext not in PPT_EXTENSIONS:
        return {
            "success": False,
            "source_file": ppt_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：不支持的文件格式 '{ext}'，仅支持 .ppt 和 .pptx"
        }
    
    # 获取幻灯片数量
    slide_count = None
    try:
        from pptx import Presentation
        prs = Presentation(ppt_path)
        slide_count = len(prs.slides)
    except:
        pass
    
    # 确定输出路径
    if output_path is None:
        base_name = os.path.splitext(os.path.basename(ppt_path))[0]
        output_path = os.path.join(
            tempfile.gettempdir(),
            f"{base_name}_{uuid.uuid4().hex[:8]}.pdf"
        )
    
    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 选择转换器
    actual_converter = converter
    if converter == 'auto':
        actual_converter = _detect_best_converter()
    
    if actual_converter == 'none':
        return {
            "success": False,
            "source_file": ppt_path,
            "pdf_path": None,
            "download_url": None,
            "message": "错误：未找到可用的转换器"
        }
    
    # 执行转换
    if actual_converter == 'libreoffice':
        success, error = _convert_with_libreoffice(ppt_path, output_path)
    elif actual_converter == 'comtypes':
        success, error = _convert_with_comtypes(ppt_path, output_path)
    elif actual_converter == 'python':
        success, error = _convert_with_python(ppt_path, output_path)
    else:
        return {
            "success": False,
            "source_file": ppt_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：未知的转换器 '{actual_converter}'"
        }
    
    if not success:
        return {
            "success": False,
            "source_file": ppt_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"转换失败: {error}",
            "converter_used": actual_converter
        }
    
    if not os.path.exists(output_path):
        return {
            "success": False,
            "source_file": ppt_path,
            "pdf_path": None,
            "download_url": None,
            "message": "错误：转换失败，未生成 PDF 文件",
            "converter_used": actual_converter
        }
    
    download_url = register_file_download(output_path)
    
    result = {
        "success": True,
        "source_file": ppt_path,
        "pdf_path": output_path,
        "download_url": download_url,
        "message": f"转换成功：{os.path.basename(ppt_path)} -> {os.path.basename(output_path)}",
        "converter_used": actual_converter,
        "file_size_kb": round(os.path.getsize(output_path) / 1024, 2)
    }
    
    if slide_count is not None:
        result["slide_count"] = slide_count
    
    return result


@tool
def ppt_to_pdf(
    input_path: str,
    converter: str = "auto",
    output_dir: Optional[str] = None,
) -> dict:
    """将 PPT/PPTX 文件转换为 PDF 格式。
    
    支持单个或多个文件输入，每个 PPT 生成一个对应的 PDF（一对一）。

    Args:
        input_path: 输入文件路径，支持：
            - 单个文件："/path/to/file.pptx"
            - 多个文件（逗号分隔）："/path/file1.pptx, /path/file2.pptx"
        converter: 转换器，可选 'libreoffice'/'comtypes'/'python'/'auto'（默认自动检测）。
        output_dir: 输出目录，若不指定则使用临时目录。

    Returns:
        包含转换结果的字典：
        - success: 是否成功
        - total_files: 处理的文件总数
        - processed: 成功数
        - failed: 失败数
        - results: 每个文件的转换结果（含 source_file, pdf_path, download_url）
        - file_mappings: 文件对应关系列表
        - message: 状态信息
        - converter_used: 使用的转换器
        
        单文件时额外返回：pdf_path, download_url, slide_count, file_size_kb
    """
    if not input_path or not input_path.strip():
        return {
            "success": False,
            "total_files": 0,
            "processed": 0,
            "failed": 0,
            "results": [],
            "message": "错误：请提供有效的输入路径"
        }
    
    # 解析输入路径（支持逗号分隔）
    input_paths = [p.strip() for p in input_path.split(',') if p.strip()]
    
    # 验证并收集所有 PPT 文件
    all_ppt_files = []
    invalid_paths = []
    for path in input_paths:
        valid_path = _validate_ppt_file(path)
        if valid_path:
            all_ppt_files.append(valid_path)
        else:
            invalid_paths.append(path)
    
    all_ppt_files = sorted(set(all_ppt_files))
    
    if not all_ppt_files:
        error_msg = f"错误：未找到有效的 PPT 文件（输入：{input_path}）"
        if invalid_paths:
            error_msg += f"\n无效路径: {', '.join(invalid_paths)}"
        return {
            "success": False,
            "total_files": 0,
            "processed": 0,
            "failed": 0,
            "results": [],
            "message": error_msg
        }
    
    # 确定输出目录
    if output_dir is None:
        output_dir = os.path.join(tempfile.gettempdir(), f"ppt2pdf_{uuid.uuid4().hex[:8]}")
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 检测转换器（只检测一次）
    actual_converter = converter
    if converter == 'auto':
        actual_converter = _detect_best_converter()
    
    # 处理所有文件
    results = []
    processed = 0
    failed = 0
    
    # 多文件 + LibreOffice 时使用批量转换（大幅提升速度）
    if len(all_ppt_files) > 1 and actual_converter == 'libreoffice':
        # 批量转换：只启动一次 LibreOffice
        batch_results = _convert_batch_with_libreoffice(all_ppt_files, output_dir)
        
        for ppt_file in all_ppt_files:
            success, result_or_error = batch_results.get(ppt_file, (False, "未知错误"))
            
            # 获取幻灯片数量
            slide_count = None
            try:
                from pptx import Presentation as PptxPresentation
                prs = PptxPresentation(ppt_file)
                slide_count = len(prs.slides)
            except:
                pass
            
            if success:
                pdf_path = result_or_error
                download_url = register_file_download(pdf_path)
                
                result = {
                    "success": True,
                    "source_file": ppt_file,
                    "pdf_path": pdf_path,
                    "download_url": download_url,
                    "message": f"转换成功：{os.path.basename(ppt_file)} -> {os.path.basename(pdf_path)}",
                    "converter_used": actual_converter,
                    "file_size_kb": round(os.path.getsize(pdf_path) / 1024, 2)
                }
                if slide_count is not None:
                    result["slide_count"] = slide_count
                processed += 1
            else:
                result = {
                    "success": False,
                    "source_file": ppt_file,
                    "pdf_path": None,
                    "download_url": None,
                    "message": result_or_error,
                    "converter_used": actual_converter
                }
                failed += 1
            
            results.append(result)
    else:
        # 单文件或非 LibreOffice：逐个处理
        for ppt_file in all_ppt_files:
            base_name = os.path.splitext(os.path.basename(ppt_file))[0]
            output_path = os.path.join(output_dir, f"{base_name}.pdf")
            
            # 避免文件名冲突
            counter = 1
            while os.path.exists(output_path):
                output_path = os.path.join(output_dir, f"{base_name}_{counter}.pdf")
                counter += 1
            
            result = _convert_single_ppt(
                ppt_path=ppt_file,
                output_path=output_path,
                converter=actual_converter,
            )
            
            results.append(result)
            
            if result["success"]:
                processed += 1
            else:
                failed += 1
    
    # 构建返回结果
    success = failed == 0 and processed > 0
    
    # 构建文件映射列表
    file_mappings = []
    for r in results:
        if r["success"]:
            file_mappings.append({
                "source": r["source_file"],
                "output": r["pdf_path"],
                "download_url": r["download_url"],
            })
    
    response = {
        "success": success,
        "total_files": len(all_ppt_files),
        "processed": processed,
        "failed": failed,
        "results": results,
        "file_mappings": file_mappings,
        "output_dir": output_dir,
        "converter_used": actual_converter,
        "message": f"转换完成：共 {len(all_ppt_files)} 个文件，成功 {processed} 个，失败 {failed} 个"
    }
    
    # 单文件时返回简化结果
    if len(all_ppt_files) == 1:
        single = results[0]
        response.update({
            "pdf_path": single.get("pdf_path"),
            "download_url": single.get("download_url"),
            "slide_count": single.get("slide_count"),
            "file_size_kb": single.get("file_size_kb"),
            "message": single.get("message"),
        })
    
    return response


if __name__ == "__main__":
    # 检测环境
    print("=== 检测转换器 ===")
    converter = _detect_best_converter()
    print(f"可用转换器: {converter}")
    
    # 测试
    test_dir = r"D:\test"
    if os.path.exists(test_dir):
        ppt_files = [f for f in os.listdir(test_dir) if f.lower().endswith(PPT_EXTENSIONS)]
        if ppt_files:
            print(f"\n=== 测试单文件 ===")
            test_file = os.path.join(test_dir, ppt_files[0])
            result = ppt_to_pdf.invoke({"input_path": test_file})
            print(f"成功: {result['success']}")
            print(f"消息: {result['message']}")
            print(f"下载链接: {result.get('download_url')}")
            
            if len(ppt_files) >= 2:
                print(f"\n=== 测试多文件 ===")
                multi_paths = ", ".join([os.path.join(test_dir, f) for f in ppt_files[:2]])
                result = ppt_to_pdf.invoke({"input_path": multi_paths})
                print(f"成功: {result['success']}")
                print(f"消息: {result['message']}")
                print("文件对应关系:")
                for mapping in result.get('file_mappings', []):
                    print(f"  {os.path.basename(mapping['source'])} -> {os.path.basename(mapping['output'])}")
    else:
        print(f"测试目录不存在: {test_dir}")
        print("工具已加载成功，可上传到平台使用")
