"""
PPT/PPTX 转 PDF 工具

将 PowerPoint 演示文稿转换为 PDF 格式。
支持多种转换后端：
1. LibreOffice (Linux/Docker 推荐，免费开源)
2. comtypes + PowerPoint (Windows，需要安装 Microsoft PowerPoint)
3. python-pptx + reportlab (纯Python方案，基础转换)

安装依赖：
- Linux/Docker: apt-get install libreoffice
- Windows: pip install comtypes (需要安装 Microsoft PowerPoint)
- 纯Python: pip install python-pptx reportlab Pillow
"""
import os
import subprocess
import base64
import requests
import tempfile
import uuid
import shutil
from typing import Optional, List
from langchain_core.tools import tool

# MCP 服务地址（用于注册文件下载）
MCP_API_BASE = os.environ.get("MCP_API_BASE", "http://localhost:7000/api/mcp")


def register_file_download(file_path: str, expires_in: int = 3600) -> Optional[str]:
    """向 MCP 服务注册文件，获取下载 URL"""
    try:
        resp = requests.post(
            f"{MCP_API_BASE}/files/register",
            json={"file_path": file_path, "expires_in": expires_in},
            timeout=10
        )
        if resp.status_code == 200:
            return resp.json().get("download_url")
    except Exception:
        pass
    return None


def _convert_with_libreoffice(ppt_path: str, output_path: str) -> tuple:
    """使用 LibreOffice 转换（Linux/Docker 推荐）"""
    try:
        # 获取输出目录
        output_dir = os.path.dirname(output_path) or tempfile.gettempdir()
        
        # 使用 LibreOffice 命令行转换
        cmd = [
            'soffice',
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
            ppt_path
        ]
        
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=180  # PPT转换可能较慢
        )
        
        if result.returncode != 0:
            return False, f"LibreOffice 转换失败: {result.stderr}"
        
        # LibreOffice 输出的文件名
        base_name = os.path.splitext(os.path.basename(ppt_path))[0]
        generated_pdf = os.path.join(output_dir, f"{base_name}.pdf")
        
        # 如果需要重命名到指定路径
        if generated_pdf != output_path and os.path.exists(generated_pdf):
            shutil.move(generated_pdf, output_path)
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except FileNotFoundError:
        return False, "未找到 LibreOffice，请安装: apt-get install libreoffice"
    except subprocess.TimeoutExpired:
        return False, "转换超时，PPT文件可能过大或过于复杂"
    except Exception as e:
        return False, f"LibreOffice 转换异常: {str(e)}"


def _convert_with_comtypes(ppt_path: str, output_path: str) -> tuple:
    """使用 comtypes + PowerPoint 转换（Windows）"""
    try:
        import comtypes.client
        
        # 初始化 PowerPoint
        powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
        powerpoint.Visible = 1
        
        try:
            # 打开演示文稿
            presentation = powerpoint.Presentations.Open(ppt_path)
            
            # 保存为 PDF (formatType=32 表示 PDF)
            presentation.SaveAs(output_path, 32)
            
            # 关闭演示文稿
            presentation.Close()
            
        finally:
            # 退出 PowerPoint
            powerpoint.Quit()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except ImportError:
        return False, "未安装 comtypes，请执行: pip install comtypes"
    except Exception as e:
        error_msg = str(e)
        if "Powerpoint.Application" in error_msg:
            return False, "未安装 Microsoft PowerPoint 或无法启动"
        return False, f"comtypes 转换异常: {error_msg}"


def _convert_with_python(ppt_path: str, output_path: str) -> tuple:
    """使用纯Python库转换（基础转换，主要提取文本和图片）"""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from PIL import Image
        import io
        
        # 尝试注册中文字体
        font_name = 'Helvetica'
        try:
            chinese_fonts = [
                '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
                '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
                'C:/Windows/Fonts/simhei.ttf',
                'C:/Windows/Fonts/msyh.ttc',
            ]
            for font_path in chinese_fonts:
                if os.path.exists(font_path):
                    pdfmetrics.registerFont(TTFont('ChineseFont', font_path))
                    font_name = 'ChineseFont'
                    break
        except:
            pass
        
        # 读取 PPT
        prs = Presentation(ppt_path)
        
        # 获取幻灯片尺寸
        slide_width = prs.slide_width.inches
        slide_height = prs.slide_height.inches
        
        # 创建 PDF (使用幻灯片的实际尺寸)
        page_size = (slide_width * 72, slide_height * 72)  # 转换为点
        c = canvas.Canvas(output_path, pagesize=page_size)
        
        for slide_num, slide in enumerate(prs.slides, 1):
            # 设置页面背景为白色
            c.setFillColorRGB(1, 1, 1)
            c.rect(0, 0, page_size[0], page_size[1], fill=1)
            
            # 提取文本
            y_position = page_size[1] - 50
            
            for shape in slide.shapes:
                # 处理文本框
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            try:
                                c.setFont(font_name, 12)
                                c.setFillColorRGB(0, 0, 0)
                                
                                # 处理长文本换行
                                max_width = page_size[0] - 100
                                if c.stringWidth(text, font_name, 12) > max_width:
                                    # 简单截断处理
                                    while c.stringWidth(text + "...", font_name, 12) > max_width and len(text) > 10:
                                        text = text[:-1]
                                    text = text + "..."
                                
                                c.drawString(50, y_position, text)
                                y_position -= 20
                                
                                if y_position < 50:
                                    break
                            except:
                                pass
                
                # 处理图片
                if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        
                        # 使用 PIL 处理图片
                        img = Image.open(io.BytesIO(image_bytes))
                        img_reader = ImageReader(img)
                        
                        # 计算图片位置和大小
                        img_width = min(shape.width.inches * 72, page_size[0] - 100)
                        img_height = min(shape.height.inches * 72, page_size[1] - 100)
                        
                        x_pos = (page_size[0] - img_width) / 2
                        y_pos = max(50, y_position - img_height)
                        
                        c.drawImage(img_reader, x_pos, y_pos, 
                                   width=img_width, height=img_height,
                                   preserveAspectRatio=True)
                        
                        y_position = y_pos - 20
                    except:
                        pass
            
            # 添加页码
            c.setFont(font_name, 10)
            c.setFillColorRGB(0.5, 0.5, 0.5)
            c.drawCentredString(page_size[0] / 2, 20, f"第 {slide_num} 页")
            
            # 新建下一页
            c.showPage()
        
        c.save()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except ImportError as e:
        return False, f"缺少依赖库: {str(e)}，请执行: pip install python-pptx reportlab Pillow"
    except Exception as e:
        return False, f"Python 转换异常: {str(e)}"


def _detect_best_converter() -> str:
    """检测最佳的转换方案"""
    # 优先检查 LibreOffice（Linux/Docker 环境推荐）
    try:
        result = subprocess.run(['soffice', '--version'], capture_output=True, timeout=5)
        if result.returncode == 0:
            return 'libreoffice'
    except:
        pass
    
    # 检查 comtypes（Windows 环境）
    try:
        import comtypes.client
        # 尝试检测 PowerPoint 是否可用
        try:
            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.Quit()
            return 'comtypes'
        except:
            pass
    except ImportError:
        pass
    
    # 最后使用纯 Python 方案
    try:
        import pptx
        import reportlab
        return 'python'
    except ImportError:
        pass
    
    return 'none'


@tool
def ppt_to_pdf(
    ppt_path: str,
    output_path: Optional[str] = None,
    return_base64: bool = False,
    converter: Optional[str] = None
) -> dict:
    """将 PPT/PPTX 文件转换为 PDF 格式

    Args:
        ppt_path: PowerPoint 文件的完整路径，支持 .ppt 和 .pptx 格式
        output_path: PDF 输出路径，默认在临时目录生成
        return_base64: 是否返回 PDF 的 Base64 编码内容，默认 False
        converter: 指定转换器，可选 'libreoffice'/'comtypes'/'python'/'auto'(默认自动检测)

    Returns:
        包含转换结果的字典：
        - success: 是否成功
        - pdf_path: 生成的 PDF 文件路径（服务器本地）
        - download_url: 远程可访问的下载链接
        - message: 结果消息
        - converter_used: 使用的转换器
        - file_size_kb: 文件大小(KB)
        - slide_count: 幻灯片数量（如果可获取）
        - base64_content: (可选) PDF 的 Base64 编码内容
    """
    try:
        # 检查源文件是否存在
        if not os.path.exists(ppt_path):
            return {
                "success": False,
                "pdf_path": None,
                "download_url": None,
                "message": f"错误：文件不存在 - {ppt_path}"
            }

        # 检查文件扩展名
        ext = os.path.splitext(ppt_path)[1].lower()
        if ext not in ['.ppt', '.pptx']:
            return {
                "success": False,
                "pdf_path": None,
                "download_url": None,
                "message": f"错误：不支持的文件格式 '{ext}'，仅支持 .ppt 和 .pptx"
            }

        # 尝试获取幻灯片数量
        slide_count = None
        try:
            from pptx import Presentation
            prs = Presentation(ppt_path)
            slide_count = len(prs.slides)
        except:
            pass

        # 确定输出路径
        if output_path is None:
            base_name = os.path.splitext(os.path.basename(ppt_path))[0]
            output_path = os.path.join(
                tempfile.gettempdir(),
                f"{base_name}_{uuid.uuid4().hex[:8]}.pdf"
            )

        # 确保输出目录存在
        output_dir = os.path.dirname(output_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)

        # 选择转换器
        if converter is None or converter == 'auto':
            converter = _detect_best_converter()
        
        if converter == 'none':
            return {
                "success": False,
                "pdf_path": None,
                "download_url": None,
                "message": "错误：未找到可用的转换器。请安装依赖：\n"
                          "- Linux/Docker: apt-get install libreoffice\n"
                          "- Windows: pip install comtypes (需要 Microsoft PowerPoint)\n"
                          "- 纯Python: pip install python-pptx reportlab Pillow"
            }

        # 执行转换
        if converter == 'libreoffice':
            success, error = _convert_with_libreoffice(ppt_path, output_path)
        elif converter == 'comtypes':
            success, error = _convert_with_comtypes(ppt_path, output_path)
        elif converter == 'python':
            success, error = _convert_with_python(ppt_path, output_path)
        else:
            return {
                "success": False,
                "pdf_path": None,
                "download_url": None,
                "message": f"错误：未知的转换器 '{converter}'，可选: libreoffice/comtypes/python/auto"
            }

        if not success:
            return {
                "success": False,
                "pdf_path": None,
                "download_url": None,
                "message": f"转换失败: {error}",
                "converter_used": converter
            }

        # 检查转换结果
        if not os.path.exists(output_path):
            return {
                "success": False,
                "pdf_path": None,
                "download_url": None,
                "message": "错误：转换失败，未生成 PDF 文件",
                "converter_used": converter
            }

        # 注册文件获取下载 URL
        download_url = register_file_download(output_path)

        result = {
            "success": True,
            "pdf_path": output_path,
            "download_url": download_url,
            "message": f"转换成功：{os.path.basename(ppt_path)} -> {os.path.basename(output_path)}",
            "converter_used": converter,
            "file_size_kb": round(os.path.getsize(output_path) / 1024, 2)
        }
        
        if slide_count is not None:
            result["slide_count"] = slide_count

        # 如果需要返回 Base64 内容
        if return_base64:
            with open(output_path, 'rb') as f:
                pdf_bytes = f.read()
            result["base64_content"] = base64.b64encode(pdf_bytes).decode('utf-8')

        return result

    except Exception as e:
        return {
            "success": False,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：转换过程中出现异常 - {str(e)}"
        }


@tool
def batch_ppt_to_pdf(
    ppt_paths: list,
    output_dir: Optional[str] = None,
    converter: Optional[str] = None
) -> dict:
    """批量将多个 PPT/PPTX 文件转换为 PDF 格式

    Args:
        ppt_paths: PowerPoint 文件路径列表
        output_dir: PDF 输出目录，默认使用临时目录
        converter: 指定转换器，可选 'libreoffice'/'comtypes'/'python'/'auto'(默认)

    Returns:
        包含批量转换结果的字典：
        - success: 是否全部成功
        - total: 总文件数
        - converted: 成功转换数
        - failed: 失败数
        - converter_used: 使用的转换器
        - results: 各文件转换详情列表（含 download_url）
    """
    results = []
    converted = 0
    failed = 0
    
    # 自动检测转换器（只检测一次）
    if converter is None or converter == 'auto':
        converter = _detect_best_converter()

    for ppt_path in ppt_paths:
        # 确定输出路径
        if output_dir:
            filename = os.path.splitext(os.path.basename(ppt_path))[0] + '.pdf'
            output_path = os.path.join(output_dir, filename)
        else:
            output_path = None

        # 调用单文件转换
        result = ppt_to_pdf.invoke({
            "ppt_path": ppt_path,
            "output_path": output_path,
            "return_base64": False,
            "converter": converter
        })

        results.append({
            "source": ppt_path,
            "result": result
        })

        if result["success"]:
            converted += 1
        else:
            failed += 1

    return {
        "success": failed == 0,
        "total": len(ppt_paths),
        "converted": converted,
        "failed": failed,
        "converter_used": converter,
        "results": results
    }


@tool
def check_ppt2pdf_environment() -> dict:
    """检查 PPT 转 PDF 的运行环境和可用转换器

    Returns:
        包含环境信息的字典：
        - available_converters: 可用的转换器列表
        - recommended: 推荐使用的转换器
        - versions: 各组件版本信息
        - message: 环境说明
    """
    available = []
    versions = {}
    
    # 检查 LibreOffice
    try:
        result = subprocess.run(
            ['soffice', '--version'], 
            capture_output=True, 
            text=True, 
            timeout=5
        )
        if result.returncode == 0:
            available.append('libreoffice')
            versions['libreoffice'] = result.stdout.strip()
    except:
        pass
    
    # 检查 comtypes
    try:
        import comtypes.client
        available.append('comtypes')
        versions['comtypes'] = 'installed'
        
        # 检查 PowerPoint 是否可用
        try:
            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.Quit()
            versions['powerpoint'] = 'available'
        except:
            versions['powerpoint'] = 'not available'
    except ImportError:
        pass
    
    # 检查纯 Python 方案
    python_available = True
    try:
        import pptx
        versions['python-pptx'] = getattr(pptx, '__version__', 'installed')
    except ImportError:
        python_available = False
    
    try:
        import reportlab
        versions['reportlab'] = getattr(reportlab, 'Version', 'installed')
    except ImportError:
        python_available = False
    
    try:
        from PIL import Image
        import PIL
        versions['pillow'] = PIL.__version__
    except ImportError:
        python_available = False
    
    if python_available:
        available.append('python')
    
    # 确定推荐
    if 'libreoffice' in available:
        recommended = 'libreoffice'
        message = "LibreOffice 可用，推荐在 Linux/Docker 环境下使用，转换质量最佳"
    elif 'comtypes' in available and versions.get('powerpoint') == 'available':
        recommended = 'comtypes'
        message = "PowerPoint 可用，适合 Windows 环境，转换质量与原版一致"
    elif 'python' in available:
        recommended = 'python'
        message = "使用纯 Python 方案，可进行基础转换，复杂排版可能有差异"
    else:
        recommended = None
        message = "未找到可用的转换器，请安装依赖：\n" \
                  "- Linux/Docker: apt-get install libreoffice\n" \
                  "- Windows: pip install comtypes (需要 Microsoft PowerPoint)\n" \
                  "- 纯Python: pip install python-pptx reportlab Pillow"
    
    return {
        "available_converters": available,
        "recommended": recommended,
        "versions": versions,
        "message": message
    }


# 测试代码
if __name__ == "__main__":
    # 检查环境
    print("=== 检查运行环境 ===")
    env_result = check_ppt2pdf_environment.invoke({})
    print(f"可用转换器: {env_result['available_converters']}")
    print(f"推荐使用: {env_result['recommended']}")
    print(f"版本信息: {env_result['versions']}")
    print(f"说明: {env_result['message']}")
    print()
    
    # 单文件测试
    test_ppt = r"D:\test\sample.pptx"
    if os.path.exists(test_ppt):
        print("=== 单文件转换测试 ===")
        result = ppt_to_pdf.invoke({"ppt_path": test_ppt})
        print(result)
    else:
        print(f"测试文件不存在: {test_ppt}")
        print("工具已加载成功，可上传到平台使用")
