"""Word 文档转 PDF 工具 - 支持单文件和多文件输入

将 Word 文档（.doc/.docx）转换为 PDF 格式。
支持多种转换后端：
1. LibreOffice (Linux/Docker 推荐，免费开源)
2. comtypes + Word (Windows，需要安装 Microsoft Word)
3. python-docx + reportlab (纯Python方案，基础转换)
"""
import os
import subprocess
import tempfile
import uuid
import shutil
import requests
from typing import Optional, List, Dict

from langchain_core.tools import tool

# MCP 服务地址（用于注册文件下载）
MCP_API_BASE = os.environ.get("MCP_API_BASE", "http://localhost:7000/api/mcp")

# 支持的 Word 文件扩展名
WORD_EXTENSIONS = ('.doc', '.docx')


def register_file_download(file_path: str, expires_in: int = 3600) -> Optional[str]:
    """向 MCP 服务注册文件，获取下载 URL"""
    try:
        resp = requests.post(
            f"{MCP_API_BASE}/files/register",
            json={"file_path": file_path, "expires_in": expires_in},
            timeout=10
        )
        if resp.status_code == 200:
            return resp.json().get("download_url")
    except Exception:
        pass
    return None


def _validate_word_file(path: str) -> Optional[str]:
    """验证并返回有效的 Word 文件路径
    
    Args:
        path: 文件路径
        
    Returns:
        有效的绝对路径，如果文件无效则返回 None
    """
    if os.path.isfile(path) and path.lower().endswith(WORD_EXTENSIONS):
        return os.path.abspath(path)
    return None


def _convert_with_libreoffice(doc_path: str, output_path: str) -> tuple:
    """使用 LibreOffice 转换单个文件（Linux/Docker 推荐）"""
    try:
        output_dir = os.path.dirname(output_path) or tempfile.gettempdir()
        
        cmd = [
            'soffice',
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
            doc_path
        ]
        
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=180
        )
        
        if result.returncode != 0:
            return False, f"LibreOffice 转换失败: {result.stderr}"
        
        base_name = os.path.splitext(os.path.basename(doc_path))[0]
        generated_pdf = os.path.join(output_dir, f"{base_name}.pdf")
        
        if generated_pdf != output_path and os.path.exists(generated_pdf):
            shutil.move(generated_pdf, output_path)
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except FileNotFoundError:
        return False, "未找到 LibreOffice，请安装: apt-get install libreoffice"
    except subprocess.TimeoutExpired:
        return False, "转换超时，文档可能过大或过于复杂"
    except Exception as e:
        return False, f"LibreOffice 转换异常: {str(e)}"


def _convert_batch_with_libreoffice(doc_files: List[str], output_dir: str) -> Dict[str, tuple]:
    """使用 LibreOffice 批量转换多个文件（只启动一次，大幅提升速度）
    
    Args:
        doc_files: Word 文件路径列表
        output_dir: 输出目录
        
    Returns:
        字典 {doc_path: (success, error_or_pdf_path)}
    """
    results = {}
    
    if not doc_files:
        return results
    
    try:
        # 批量转换：一次性传入所有文件，只启动一次 LibreOffice
        cmd = [
            'soffice',
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
        ] + doc_files
        
        # 根据文件数量动态调整超时时间（每个文件约60秒）
        timeout = max(180, len(doc_files) * 60)
        
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=timeout
        )
        
        # 检查每个文件的转换结果
        for doc_path in doc_files:
            base_name = os.path.splitext(os.path.basename(doc_path))[0]
            generated_pdf = os.path.join(output_dir, f"{base_name}.pdf")
            
            if os.path.exists(generated_pdf):
                results[doc_path] = (True, generated_pdf)
            else:
                results[doc_path] = (False, f"转换失败，未生成 PDF（可能是文件损坏或格式不支持）")
        
        # 如果整体命令失败且没有任何成功的文件
        if result.returncode != 0 and not any(r[0] for r in results.values()):
            error_msg = f"LibreOffice 批量转换失败: {result.stderr}"
            for doc_path in doc_files:
                if doc_path not in results or not results[doc_path][0]:
                    results[doc_path] = (False, error_msg)
                    
    except FileNotFoundError:
        error = "未找到 LibreOffice，请安装: apt-get install libreoffice"
        for doc_path in doc_files:
            results[doc_path] = (False, error)
    except subprocess.TimeoutExpired:
        error = "批量转换超时，文件可能过多或过大"
        for doc_path in doc_files:
            if doc_path not in results:
                results[doc_path] = (False, error)
    except Exception as e:
        error = f"LibreOffice 批量转换异常: {str(e)}"
        for doc_path in doc_files:
            results[doc_path] = (False, error)
    
    return results


def _convert_with_comtypes(doc_path: str, output_path: str) -> tuple:
    """使用 comtypes + Word 转换（Windows）"""
    try:
        import comtypes.client
        
        word = comtypes.client.CreateObject("Word.Application")
        word.Visible = False
        
        try:
            doc = word.Documents.Open(doc_path)
            # 17 = wdFormatPDF
            doc.SaveAs(output_path, FileFormat=17)
            doc.Close()
        finally:
            word.Quit()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except ImportError:
        return False, "未安装 comtypes，请执行: pip install comtypes"
    except Exception as e:
        error_msg = str(e)
        if "Word.Application" in error_msg:
            return False, "未安装 Microsoft Word 或无法启动"
        return False, f"comtypes 转换异常: {error_msg}"


def _convert_with_python(doc_path: str, output_path: str) -> tuple:
    """使用纯Python库转换（基础转换，仅支持 .docx）"""
    try:
        from docx import Document
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.units import inch
        
        # 检查是否为 .docx 格式
        if not doc_path.lower().endswith('.docx'):
            return False, "纯 Python 方案仅支持 .docx 格式，.doc 文件请使用 LibreOffice 或 Word"
        
        # 注册中文字体
        font_name = 'Helvetica'
        try:
            chinese_fonts = [
                '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
                '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
                'C:/Windows/Fonts/simhei.ttf',
                'C:/Windows/Fonts/msyh.ttc',
            ]
            for font_path in chinese_fonts:
                if os.path.exists(font_path):
                    pdfmetrics.registerFont(TTFont('ChineseFont', font_path))
                    font_name = 'ChineseFont'
                    break
        except:
            pass
        
        # 读取 Word 文档
        doc = Document(doc_path)
        
        # 创建 PDF
        c = canvas.Canvas(output_path, pagesize=A4)
        width, height = A4
        
        y_position = height - 50
        line_height = 14
        margin = 50
        max_width = width - 2 * margin
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                y_position -= line_height
                continue
            
            # 简单的文本换行处理
            c.setFont(font_name, 11)
            
            # 检查是否需要新页
            if y_position < margin + line_height:
                c.showPage()
                y_position = height - 50
                c.setFont(font_name, 11)
            
            # 绘制文本（简单处理，不支持复杂格式）
            try:
                # 长文本分行
                words = text
                current_line = ""
                for char in words:
                    test_line = current_line + char
                    if c.stringWidth(test_line, font_name, 11) < max_width:
                        current_line = test_line
                    else:
                        if current_line:
                            c.drawString(margin, y_position, current_line)
                            y_position -= line_height
                            if y_position < margin + line_height:
                                c.showPage()
                                y_position = height - 50
                                c.setFont(font_name, 11)
                        current_line = char
                
                if current_line:
                    c.drawString(margin, y_position, current_line)
                    y_position -= line_height
            except:
                pass
            
            y_position -= 5  # 段落间距
        
        c.save()
        
        if os.path.exists(output_path):
            return True, None
        else:
            return False, "转换后未生成PDF文件"
            
    except ImportError as e:
        return False, f"缺少依赖库: {str(e)}，请执行: pip install python-docx reportlab"
    except Exception as e:
        return False, f"Python 转换异常: {str(e)}"


def _detect_best_converter() -> str:
    """检测最佳的转换方案"""
    # 优先检查 LibreOffice（Linux/Docker 环境推荐）
    try:
        result = subprocess.run(['soffice', '--version'], capture_output=True, timeout=5)
        if result.returncode == 0:
            return 'libreoffice'
    except:
        pass
    
    # 检查 comtypes（Windows 环境）
    try:
        import comtypes.client
        try:
            word = comtypes.client.CreateObject("Word.Application")
            word.Quit()
            return 'comtypes'
        except:
            pass
    except ImportError:
        pass
    
    # 最后使用纯 Python 方案
    try:
        import docx
        import reportlab
        return 'python'
    except ImportError:
        pass
    
    return 'none'


def _get_page_count(doc_path: str) -> Optional[int]:
    """获取文档页数（仅 .docx 支持估算）"""
    try:
        from docx import Document
        doc = Document(doc_path)
        # 简单估算：每页约 40 段落
        para_count = len(doc.paragraphs)
        return max(1, para_count // 40 + 1)
    except:
        return None


def _convert_single_doc(
    doc_path: str,
    output_path: Optional[str] = None,
    converter: str = "auto",
) -> dict:
    """转换单个 Word 文件为 PDF（内部函数）"""
    
    if not os.path.isfile(doc_path):
        return {
            "success": False,
            "source_file": doc_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：文件不存在 - {doc_path}"
        }
    
    ext = os.path.splitext(doc_path)[1].lower()
    if ext not in WORD_EXTENSIONS:
        return {
            "success": False,
            "source_file": doc_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：不支持的文件格式 '{ext}'，仅支持 .doc 和 .docx"
        }
    
    # 获取页数估算
    page_count = _get_page_count(doc_path)
    
    # 确定输出路径
    if output_path is None:
        base_name = os.path.splitext(os.path.basename(doc_path))[0]
        output_path = os.path.join(
            tempfile.gettempdir(),
            f"{base_name}_{uuid.uuid4().hex[:8]}.pdf"
        )
    
    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 选择转换器
    actual_converter = converter
    if converter == 'auto':
        actual_converter = _detect_best_converter()
    
    if actual_converter == 'none':
        return {
            "success": False,
            "source_file": doc_path,
            "pdf_path": None,
            "download_url": None,
            "message": "错误：未找到可用的转换器"
        }
    
    # 执行转换
    if actual_converter == 'libreoffice':
        success, error = _convert_with_libreoffice(doc_path, output_path)
    elif actual_converter == 'comtypes':
        success, error = _convert_with_comtypes(doc_path, output_path)
    elif actual_converter == 'python':
        success, error = _convert_with_python(doc_path, output_path)
    else:
        return {
            "success": False,
            "source_file": doc_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"错误：未知的转换器 '{actual_converter}'"
        }
    
    if not success:
        return {
            "success": False,
            "source_file": doc_path,
            "pdf_path": None,
            "download_url": None,
            "message": f"转换失败: {error}",
            "converter_used": actual_converter
        }
    
    if not os.path.exists(output_path):
        return {
            "success": False,
            "source_file": doc_path,
            "pdf_path": None,
            "download_url": None,
            "message": "错误：转换失败，未生成 PDF 文件",
            "converter_used": actual_converter
        }
    
    download_url = register_file_download(output_path)
    
    result = {
        "success": True,
        "source_file": doc_path,
        "pdf_path": output_path,
        "download_url": download_url,
        "message": f"转换成功：{os.path.basename(doc_path)} -> {os.path.basename(output_path)}",
        "converter_used": actual_converter,
        "file_size_kb": round(os.path.getsize(output_path) / 1024, 2)
    }
    
    if page_count is not None:
        result["page_count"] = page_count
    
    return result


@tool
def word_to_pdf(
    input_path: str,
    converter: str = "auto",
    output_dir: Optional[str] = None,
) -> dict:
    """将 Word 文档转换为 PDF 格式。
    
    支持单个或多个文件输入，每个 Word 文档生成一个对应的 PDF（一对一）。

    Args:
        input_path: 输入文件路径，支持：
            - 单个文件："/path/to/file.docx"
            - 多个文件（逗号分隔）："/path/file1.docx, /path/file2.doc"
        converter: 转换器，可选 'libreoffice'/'comtypes'/'python'/'auto'（默认自动检测）。
        output_dir: 输出目录，若不指定则使用临时目录。

    Returns:
        包含转换结果的字典：
        - success: 是否成功
        - total_files: 处理的文件总数
        - processed: 成功数
        - failed: 失败数
        - results: 每个文件的转换结果（含 source_file, pdf_path, download_url）
        - file_mappings: 文件对应关系列表
        - message: 状态信息
        - converter_used: 使用的转换器
        
        单文件时额外返回：pdf_path, download_url, page_count, file_size_kb
    """
    if not input_path or not input_path.strip():
        return {
            "success": False,
            "total_files": 0,
            "processed": 0,
            "failed": 0,
            "results": [],
            "message": "错误：请提供有效的输入路径"
        }
    
    # 解析输入路径（支持逗号分隔）
    input_paths = [p.strip() for p in input_path.split(',') if p.strip()]
    
    # 验证并收集所有 Word 文件
    all_doc_files = []
    invalid_paths = []
    for path in input_paths:
        valid_path = _validate_word_file(path)
        if valid_path:
            all_doc_files.append(valid_path)
        else:
            invalid_paths.append(path)
    
    all_doc_files = sorted(set(all_doc_files))
    
    if not all_doc_files:
        error_msg = f"错误：未找到有效的 Word 文件（输入：{input_path}）"
        if invalid_paths:
            error_msg += f"\n无效路径: {', '.join(invalid_paths)}"
        return {
            "success": False,
            "total_files": 0,
            "processed": 0,
            "failed": 0,
            "results": [],
            "message": error_msg
        }
    
    # 确定输出目录
    if output_dir is None:
        output_dir = os.path.join(tempfile.gettempdir(), f"word2pdf_{uuid.uuid4().hex[:8]}")
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 检测转换器（只检测一次）
    actual_converter = converter
    if converter == 'auto':
        actual_converter = _detect_best_converter()
    
    # 处理所有文件
    results = []
    processed = 0
    failed = 0
    
    # 多文件 + LibreOffice 时使用批量转换（大幅提升速度）
    if len(all_doc_files) > 1 and actual_converter == 'libreoffice':
        # 批量转换：只启动一次 LibreOffice
        batch_results = _convert_batch_with_libreoffice(all_doc_files, output_dir)
        
        for doc_file in all_doc_files:
            success, result_or_error = batch_results.get(doc_file, (False, "未知错误"))
            
            # 获取页数
            page_count = _get_page_count(doc_file)
            
            if success:
                pdf_path = result_or_error
                download_url = register_file_download(pdf_path)
                
                result = {
                    "success": True,
                    "source_file": doc_file,
                    "pdf_path": pdf_path,
                    "download_url": download_url,
                    "message": f"转换成功：{os.path.basename(doc_file)} -> {os.path.basename(pdf_path)}",
                    "converter_used": actual_converter,
                    "file_size_kb": round(os.path.getsize(pdf_path) / 1024, 2)
                }
                if page_count is not None:
                    result["page_count"] = page_count
                processed += 1
            else:
                result = {
                    "success": False,
                    "source_file": doc_file,
                    "pdf_path": None,
                    "download_url": None,
                    "message": result_or_error,
                    "converter_used": actual_converter
                }
                failed += 1
            
            results.append(result)
    else:
        # 单文件或非 LibreOffice：逐个处理
        for doc_file in all_doc_files:
            base_name = os.path.splitext(os.path.basename(doc_file))[0]
            output_path = os.path.join(output_dir, f"{base_name}.pdf")
            
            # 避免文件名冲突
            counter = 1
            while os.path.exists(output_path):
                output_path = os.path.join(output_dir, f"{base_name}_{counter}.pdf")
                counter += 1
            
            result = _convert_single_doc(
                doc_path=doc_file,
                output_path=output_path,
                converter=actual_converter,
            )
            
            results.append(result)
            
            if result["success"]:
                processed += 1
            else:
                failed += 1
    
    # 构建返回结果
    success = failed == 0 and processed > 0
    
    # 构建文件映射列表
    file_mappings = []
    for r in results:
        if r["success"]:
            file_mappings.append({
                "source": r["source_file"],
                "output": r["pdf_path"],
                "download_url": r["download_url"],
            })
    
    response = {
        "success": success,
        "total_files": len(all_doc_files),
        "processed": processed,
        "failed": failed,
        "results": results,
        "file_mappings": file_mappings,
        "output_dir": output_dir,
        "converter_used": actual_converter,
        "message": f"转换完成：共 {len(all_doc_files)} 个文件，成功 {processed} 个，失败 {failed} 个"
    }
    
    # 单文件时返回简化结果
    if len(all_doc_files) == 1:
        single = results[0]
        response.update({
            "pdf_path": single.get("pdf_path"),
            "download_url": single.get("download_url"),
            "page_count": single.get("page_count"),
            "file_size_kb": single.get("file_size_kb"),
            "message": single.get("message"),
        })
    
    return response


if __name__ == "__main__":
    # 检测环境
    print("=== 检测转换器 ===")
    converter = _detect_best_converter()
    print(f"可用转换器: {converter}")
    
    # 测试
    test_dir = r"D:\test"
    if os.path.exists(test_dir):
        doc_files = [f for f in os.listdir(test_dir) if f.lower().endswith(WORD_EXTENSIONS)]
        if doc_files:
            print(f"\n=== 测试单文件 ===")
            test_file = os.path.join(test_dir, doc_files[0])
            result = word_to_pdf.invoke({"input_path": test_file})
            print(f"成功: {result['success']}")
            print(f"消息: {result['message']}")
            print(f"下载链接: {result.get('download_url')}")
            
            if len(doc_files) >= 2:
                print(f"\n=== 测试多文件 ===")
                multi_paths = ", ".join([os.path.join(test_dir, f) for f in doc_files[:2]])
                result = word_to_pdf.invoke({"input_path": multi_paths})
                print(f"成功: {result['success']}")
                print(f"消息: {result['message']}")
                print("文件对应关系:")
                for mapping in result.get('file_mappings', []):
                    print(f"  {os.path.basename(mapping['source'])} -> {os.path.basename(mapping['output'])}")
    else:
        print(f"测试目录不存在: {test_dir}")
        print("工具已加载成功，可上传到平台使用")
